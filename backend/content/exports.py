from io import BytesIO
from html import escape
from bs4 import BeautifulSoup
from django.conf import settings
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.views.decorators.http import require_GET
from docx import Document
from docx.shared import Pt
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from .models import SiteSettings, ResumeSection
from .rendering import rich_text


def plain(value):
    return BeautifulSoup(rich_text(value), "html.parser").get_text(" ", strip=True)


def resume_blocks():
    s = get_object_or_404(SiteSettings, pk=1)
    blocks = [
        ("name", s.name),
        ("text", s.title),
        ("text", " | ".join(filter(None, [s.email, settings.SITE_ORIGIN]))),
    ]
    for section in ResumeSection.objects.filter(visible=True).prefetch_related(
        "jobs__accomplishments"
    ):
        blocks.append(("section", section.title))
        if section.body:
            soup = BeautifulSoup(rich_text(section.body), "html.parser")
            for node in soup.find_all(["p", "li", "h2", "h3"]):
                blocks.append(
                    (
                        "bullet" if node.name == "li" else "text",
                        node.get_text(" ", strip=True),
                    )
                )
        for j in section.jobs.all():
            if not j.visible:
                continue
            blocks.append(("job", f"{j.title} | {j.company}"))
            end = j.end.strftime("%m/%Y") if j.end else "Present"
            date = f"{j.start:%m/%Y} - {end}"
            blocks.append(
                ("text", " | ".join(filter(None, [date, j.location, j.job_type])))
            )
            blocks.append(("text", plain(j.description)))
            blocks.extend(("bullet", a.text) for a in j.accomplishments.all())
    return blocks


@require_GET
def download(request, fmt):
    if fmt not in ["pdf", "docx"]:
        return HttpResponse(status=404)
    blocks, stream = resume_blocks(), BytesIO()
    if fmt == "docx":
        doc = Document()
        doc.styles["Normal"].font.name = "Arial"
        doc.styles["Normal"].font.size = Pt(10.5)
        for kind, text in blocks:
            if kind == "name":
                doc.add_heading(text, 0)
            elif kind in ["section", "job"]:
                doc.add_heading(text, 1 if kind == "section" else 2)
            else:
                doc.add_paragraph(
                    text, style="List Bullet" if kind == "bullet" else None
                )
        doc.save(stream)
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        fonts = settings.BASE_DIR / "fonts"
        pdfmetrics.registerFont(TTFont("Resume", str(fonts / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(
            TTFont("ResumeBold", str(fonts / "DejaVuSans-Bold.ttf"))
        )
        styles = getSampleStyleSheet()
        styles.add(
            ParagraphStyle(
                "cv", fontName="Resume", fontSize=10, leading=14, spaceAfter=5
            )
        )
        styles.add(
            ParagraphStyle(
                "cvhead",
                parent=styles["cv"],
                fontName="ResumeBold",
                fontSize=12,
                leading=16,
                spaceBefore=12,
            )
        )
        styles.add(
            ParagraphStyle("cvname", parent=styles["cvhead"], fontSize=20, leading=24)
        )
        story = []
        for kind, text in blocks:
            style = styles[
                (
                    "cvname"
                    if kind == "name"
                    else "cvhead" if kind in ["section", "job"] else "cv"
                )
                ]
            prefix = "- " if kind == "bullet" else ""
            story.append(Paragraph(escape(prefix + text), style))
        SimpleDocTemplate(
            stream,
            title="Thomas Fassih - Resume",
            author="Thomas Fassih",
            leftMargin=45,
            rightMargin=45,
            topMargin=40,
            bottomMargin=40,
        ).build(story)
        mime = "application/pdf"
    r = HttpResponse(stream.getvalue(), content_type=mime)
    r["Content-Disposition"] = f'attachment; filename="Thomas-Fassih-Resume.{fmt}"'
    r["Cache-Control"] = "no-store"
    return r


